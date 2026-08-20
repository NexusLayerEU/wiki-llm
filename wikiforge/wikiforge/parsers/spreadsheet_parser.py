"""Excel and CSV. Sheets and rows become tables, one section per sheet."""
import csv
from pathlib import Path

from .base import BaseParser, ParsedDocument, ParsedSection, ParsedTable, ParserError

#: A guard, not a preference: a 200k-row export would otherwise be pasted whole
#: into an LLM prompt. The cap is per sheet and is reported as a warning.
MAX_ROWS_PER_SHEET = 500


class XlsxParser(BaseParser):
    supported_extensions = [".xlsx", ".xlsm"]

    def parse(self, filepath: str) -> ParsedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as cause:  # pragma: no cover
            raise ParserError("openpyxl is not installed") from cause

        path = Path(filepath)
        try:
            workbook = load_workbook(str(path), read_only=True, data_only=True)
        except Exception as cause:
            raise ParserError(f"Could not open {path.name}: {cause}") from cause

        document = ParsedDocument(title=path.stem)
        text_parts: list[str] = []

        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= MAX_ROWS_PER_SHEET:
                    document.warnings.append(
                        f"Sheet '{sheet.title}' truncated at {MAX_ROWS_PER_SHEET} rows."
                    )
                    break
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    rows.append(cells)
            if not rows:
                continue
            table = ParsedTable(caption=sheet.title, headers=rows[0], rows=rows[1:])
            document.tables.append(table)
            markdown = table.to_markdown()
            text_parts.append(markdown)
            document.sections.append(
                ParsedSection(heading=sheet.title, level=1, content=markdown)
            )

        workbook.close()
        document.metadata = {"sheet_count": len(document.tables)}
        document.full_text = "\n\n".join(text_parts)
        if not document.full_text.strip():
            document.warnings.append("Every sheet was empty.")
        return document.finalise()


class CsvParser(BaseParser):
    supported_extensions = [".csv", ".tsv"]

    def parse(self, filepath: str) -> ParsedDocument:
        path = Path(filepath)
        raw = path.read_text(encoding="utf-8", errors="replace")

        # Sniff the delimiter rather than assume: .csv files are semicolon-separated
        # in much of Europe, and splitting on the wrong one yields a single column.
        try:
            dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","

        rows = [
            [cell.strip() for cell in row]
            for row in csv.reader(raw.splitlines(), delimiter=delimiter)
            if any(cell.strip() for cell in row)
        ]

        document = ParsedDocument(title=path.stem, metadata={"delimiter": delimiter})
        if not rows:
            document.warnings.append("The file had no rows.")
            return document.finalise()

        truncated = rows[: MAX_ROWS_PER_SHEET + 1]
        if len(rows) > MAX_ROWS_PER_SHEET + 1:
            document.warnings.append(f"Truncated at {MAX_ROWS_PER_SHEET} rows.")
        table = ParsedTable(caption=path.stem, headers=truncated[0], rows=truncated[1:])
        document.tables.append(table)
        document.full_text = table.to_markdown()
        document.sections.append(
            ParsedSection(heading=path.stem, level=1, content=document.full_text)
        )
        return document.finalise()
