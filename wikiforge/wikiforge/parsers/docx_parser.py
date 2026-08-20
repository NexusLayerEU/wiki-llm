"""Word documents (.docx). Legacy .doc is not readable by python-docx."""
from pathlib import Path

from .base import BaseParser, ParsedDocument, ParsedSection, ParsedTable, ParserError


class DocxParser(BaseParser):
    supported_extensions = [".docx"]

    def parse(self, filepath: str) -> ParsedDocument:
        try:
            import docx
        except ImportError as cause:  # pragma: no cover
            raise ParserError("python-docx is not installed") from cause

        path = Path(filepath)
        try:
            source = docx.Document(str(path))
        except Exception as cause:
            raise ParserError(f"Could not open {path.name}: {cause}") from cause

        document = ParsedDocument(title=path.stem)
        sections: list[ParsedSection] = []
        current = ParsedSection(heading=path.stem, level=1)
        text_parts: list[str] = []

        for paragraph in source.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower()
            if style.startswith("heading"):
                if current.content or current.heading != path.stem:
                    sections.append(current)
                level = 1
                tail = style.replace("heading", "").strip()
                if tail.isdigit():
                    level = min(int(tail), 6)
                current = ParsedSection(heading=text, level=level)
            else:
                current.content = f"{current.content}\n{text}".strip()
            text_parts.append(text)

        sections.append(current)

        for index, table in enumerate(source.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            parsed_table = ParsedTable(
                caption=f"Table {index + 1}", headers=rows[0], rows=rows[1:]
            )
            document.tables.append(parsed_table)
            text_parts.append(parsed_table.to_markdown())

        core = source.core_properties
        document.metadata = {
            "author": core.author or "",
            "created": core.created.isoformat() if core.created else "",
            "modified": core.modified.isoformat() if core.modified else "",
            "paragraph_count": len(source.paragraphs),
        }
        document.sections = [s for s in sections if s.content or s.children]
        document.full_text = "\n\n".join(text_parts)
        if not document.full_text.strip():
            document.warnings.append("No text found — the file may be scanned images.")
        return document.finalise()


class LegacyDocParser(BaseParser):
    """`.doc` is a different, binary format. Say so clearly instead of guessing."""

    supported_extensions = [".doc"]

    def parse(self, filepath: str) -> ParsedDocument:
        raise ParserError(
            "Legacy .doc is not supported. Convert it to .docx "
            "(`soffice --convert-to docx <file>`) and re-upload."
        )
